import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("python-pptx not installed.")
    sys.exit(1)

def create_3_slide_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette matching app theme
    NAVY = RGBColor(13, 27, 42)        # #0D1B2A Header/Dark
    TEAL = RGBColor(0, 168, 150)       # #00A896 Primary Accent
    LIGHT_BG = RGBColor(240, 244, 248) # #F0F4F8 Light background
    DARK_GRAY = RGBColor(45, 55, 72)   # #2D3748 Body text
    MUTED_GRAY = RGBColor(113, 128, 150) # #718096 Subtitle
    WHITE = RGBColor(255, 255, 255)

    blank_layout = prs.slide_layouts[6]

    def set_bg(slide, color):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text="PROJECT INTRODUCTION"):
        header_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        header_shape.fill.solid()
        header_shape.fill.fore_color.rgb = NAVY
        header_shape.line.color.rgb = NAVY

        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL
        p_cat.font.name = "Helvetica"

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.5), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE
        p_title.font.name = "Helvetica"

    # ==========================================
    # SLIDE 1: Title Slide (Project Introduction)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_bg(slide1, NAVY)

    # Accent decorative box
    accent = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.15), Inches(3.8))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.color.rgb = TEAL

    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.0))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "IMAGE PROCESSING & DEEP LEARNING PROJECT"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = TEAL
    p0.font.name = "Helvetica"

    p1 = tf1.add_paragraph()
    p1.text = "Butterfly Species Classifier"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.font.name = "Helvetica"

    p2 = tf1.add_paragraph()
    p2.text = "An automated system for classifying butterfly species using Neural Networks & Java Desktop GUI"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.font.name = "Helvetica"

    p3 = tf1.add_paragraph()
    p3.text = "\nClass Project Presentation"
    p3.font.size = Pt(14)
    p3.font.italic = True
    p3.font.color.rgb = MUTED_GRAY

    # ==========================================
    # SLIDE 2: What is the Project? (Simple Explanation)
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_bg(slide2, LIGHT_BG)
    add_header(slide2, "What is the Butterfly Classifier?")

    # Left Card: What it does
    card1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
    card1.fill.solid()
    card1.fill.fore_color.rgb = WHITE
    card1.line.color.rgb = RGBColor(226, 232, 240)

    tf_c1 = card1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "Overview & Purpose"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY

    bullets1 = [
        "Automated Species Identification: Recognizes 10 different species of butterflies instantly from image photos.",
        "Solves Domain Challenge: Eliminates manual identification errors for researchers and nature enthusiasts.",
        "Dataset Training: Trained on the Leeds Butterfly Dataset (832 images across 10 species classes)."
    ]
    for b in bullets1:
        p = tf_c1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY

    # Right Card: Key Features
    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.2))
    card2.fill.solid()
    card2.fill.fore_color.rgb = WHITE
    card2.line.color.rgb = RGBColor(226, 232, 240)

    tf_c2 = card2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "Key System Highlights"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEAL

    bullets2 = [
        "100% Java Desktop App: Built with Java Swing for an interactive, modern user interface.",
        "Real-Time Predictions: Displays top species name, scientific name, and confidence percentage score.",
        "Ranked Top-3 Results: Provides a ranked breakdown of top 3 species probability scores."
    ]
    for b in bullets2:
        p = tf_c2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 3: How it Works (3 Step Pipeline)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_bg(slide3, LIGHT_BG)
    add_header(slide3, "How the Project Works")

    steps = [
        ("1. Image Selection", "User selects a butterfly image (JPG/PNG) via the desktop GUI browse button.", NAVY),
        ("2. Deep Learning Inference", "Image is resized to 224x224 RGB and evaluated by MobileNetV2 Neural Network.", TEAL),
        ("3. Instant Output", "GUI displays the predicted Common Name, Scientific Name & Confidence %.", NAVY)
    ]

    for idx, (title, desc, color) in enumerate(steps):
        x = Inches(0.8 + idx * 3.9)
        card_s = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.7), Inches(5.2))
        card_s.fill.solid()
        card_s.fill.fore_color.rgb = WHITE
        card_s.line.color.rgb = RGBColor(226, 232, 240)

        tf_s = card_s.text_frame
        tf_s.word_wrap = True
        p = tf_s.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = color

        p_desc = tf_s.add_paragraph()
        p_desc.text = "\n" + desc
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = DARK_GRAY

    # Save 3-slide PPT
    output_path = "Butterfly_Classifier_3_Slide_Intro.pptx"
    prs.save(output_path)
    print(f"3-slide presentation saved to '{output_path}'.")

if __name__ == "__main__":
    create_3_slide_deck()
